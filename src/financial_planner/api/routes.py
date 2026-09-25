"""FastAPI routes for the financial planner API layer."""

import os
from io import StringIO
import io
import math
from pathlib import Path
from decimal import Decimal
from typing import List, Literal
from fastapi import APIRouter, HTTPException, UploadFile, File
from fastapi.responses import FileResponse, StreamingResponse
import pandas as pd

from src.financial_planner.api.schemas import (
    ScenarioCreate,
    ScenarioMetadataResponse,
    PriceOverrideInput,
    SimulationMetrics,
    CalculationResponse,
    CompareRequest,
    ScenarioDifference,
    LoginLogInput,
    PPTXExportInput,
)
from src.financial_planner.data_ingestion.scenario_manager import (
    list_scenarios,
    create_scenario,
    load_scenario_data,
    save_scenario_data,
    diff_scenarios_inputs,
    diff_dataframe,
    SCENARIOS_DIR,
    FILE_MAP,
)
from src.financial_planner.data_ingestion.cost_loader import load_cost_data
from src.financial_planner.data_ingestion.data_loader import load_planning_volume_data
from src.financial_planner.data_ingestion.distribution_cost_loader import load_distribution_cost_data
from src.financial_planner.data_ingestion.fx_loader import load_fx_data, load_plant_currency_data
from src.financial_planner.data_ingestion.price_loader import load_pricing_data
from src.financial_planner.data_ingestion.variable_cost_loader import load_variable_cost_data
from src.financial_planner.calculations.pricing import PriceOverride, resolve_monthly_prices
from src.financial_planner.calculations.pipeline import run_simulation_pipeline, generate_summary_metrics
from src.financial_planner.calculations.bridge import calculate_margin_bridge, summarize_margin_bridge, generate_bridge_commentary
from src.financial_planner.api.schemas import BridgeResponse
from src.financial_planner.paths import LOGIN_AUDIT_LOG

router = APIRouter()


def df_to_records(df: pd.DataFrame) -> List[dict]:
    """Helper to convert a DataFrame to JSON-serializable records."""
    if df is None:
        return []
    df_copy = df.copy()
    for col in df_copy.columns:
        # Convert Period values to strings
        df_copy[col] = df_copy[col].apply(
            lambda val: str(val) if isinstance(val, pd.Period) else val
        )
        # Convert Decimal values to float for JSON serialization
        df_copy[col] = df_copy[col].apply(
            lambda val: float(val) if isinstance(val, Decimal) else val
        )
    # Convert NaNs to None (null in JSON)
    return df_copy.where(pd.notnull(df_copy), None).to_dict(orient="records")


def add_files_status(s: dict) -> dict:
    """Helper to dynamically populate the file upload/completeness status."""
    name = s["name"]
    scen_dir = SCENARIOS_DIR / name
    files_status = {}
    for file_type, file_name in FILE_MAP.items():
        file_path = scen_dir / file_name
        if file_path.exists():
            size = file_path.stat().st_size
            if file_type == "price_overrides":
                files_status[file_type] = size > 4
            else:
                files_status[file_type] = size > 0
        else:
            files_status[file_type] = False
    s["files_status"] = files_status
    return s


@router.get("/scenarios", response_model=List[ScenarioMetadataResponse])
def get_scenarios():
    """List all available scenarios."""
    try:
        scenarios = list_scenarios()
        return [add_files_status(s) for s in scenarios]
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@router.post("/scenarios", response_model=ScenarioMetadataResponse)
def create_new_scenario(payload: ScenarioCreate):
    """Create a new scenario, optionally cloning from a base scenario."""
    try:
        s = create_scenario(payload.name, payload.base_scenario, payload.description)
        return add_files_status(s)
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@router.get("/scenarios/{name}/metadata", response_model=ScenarioMetadataResponse)
def get_scenario_metadata(name: str):
    """Fetch scenario metadata and change log."""
    scenarios = list_scenarios()
    for s in scenarios:
        if s["name"] == name:
            return add_files_status(s)
    raise HTTPException(status_code=404, detail=f"Scenario '{name}' not found.")



@router.get("/scenarios/{name}/data/{file_type}")
def get_scenario_file_data(name: str, file_type: str):
    """Fetch raw or override records for a specific scenario input file."""
    if file_type not in FILE_MAP:
        raise HTTPException(status_code=400, detail=f"Invalid file type: {file_type}")

    try:
        data = load_scenario_data(name)
        file_data = data.get(file_type)

        if file_type == "price_overrides":
            # For price_overrides, convert elements to UI format
            overrides = []
            for item in file_data:
                overrides.append({
                    "Material ID": item.get("Material ID"),
                    "Sold to ID": item.get("Sold to ID"),
                    "Ship to ID": item.get("Ship to ID"),
                    "Date": str(item.get("Date")) if item.get("Date") else None,
                    "Price": float(item.get("Price")) if isinstance(item.get("Price"), Decimal) else item.get("Price")
                })
            return overrides

        if file_data is not None:
            return df_to_records(file_data.head(100))
        return []
    except ValueError as e:
        raise HTTPException(status_code=404, detail=str(e))
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@router.post("/scenarios/{name}/overrides")
def update_price_overrides(name: str, overrides: List[PriceOverrideInput]):
    """Update and save price overrides list for a scenario."""
    try:
        # Load existing scenario data
        scen_data = load_scenario_data(name)
        
        # Format the Pydantic models to the dictionary format expected by the scenario manager
        formatted_overrides = []
        for item in overrides:
            formatted_overrides.append({
                "Material ID": item.material_id,
                "Sold to ID": item.sold_to_id,
                "Ship to ID": item.ship_to_id,
                "Date": pd.Period(item.date, freq="M"),
                "Price": item.price,
            })
        
        scen_data["price_overrides"] = formatted_overrides
        save_scenario_data(name, scen_data)
        return {"status": "success", "message": "Price overrides updated successfully."}
    except ValueError as e:
        raise HTTPException(status_code=404, detail=str(e))
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@router.post("/scenarios/{name}/upload/{file_type}")
async def upload_scenario_file(name: str, file_type: str, file: UploadFile = File(...)):
    """Upload and overwrite a scenario input CSV file after validation."""
    if file_type not in FILE_MAP or file_type == "price_overrides":
        raise HTTPException(status_code=400, detail=f"Invalid CSV file type: {file_type}")

    scen_dir = SCENARIOS_DIR / name
    if not scen_dir.exists():
        raise HTTPException(status_code=404, detail=f"Scenario '{name}' does not exist.")

    contents = await file.read()
    string_data = contents.decode("utf-8")
    io_data = StringIO(string_data)

    try:
        # Validate data by loading it through respective schema loaders
        if file_type == "volume_data":
            load_planning_volume_data(io_data)
        elif file_type == "base_prices":
            load_pricing_data(io_data)
        elif file_type == "base_costs":
            load_cost_data(io_data)
        elif file_type == "base_var_costs":
            load_variable_cost_data(io_data)
        elif file_type == "base_dist_costs":
            load_distribution_cost_data(io_data)
        elif file_type == "fx_rates":
            load_fx_data(io_data)
        elif file_type == "plant_currency":
            load_plant_currency_data(io_data)

        # If loading succeeded with no exceptions, write to disk
        dest_path = scen_dir / FILE_MAP[file_type]
        with open(dest_path, "wb") as f:
            f.write(contents)

        # Trigger auto change log recalculation by resaving the scenario
        scen_data = load_scenario_data(name)
        save_scenario_data(name, scen_data)

        return {"status": "success", "message": f"{file_type} uploaded and validated successfully."}
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Validation failed: {str(e)}")


@router.post("/scenarios/{name}/diff-file/{file_type}")
async def diff_scenario_file(name: str, file_type: str, file: UploadFile = File(...)):
    """Dry-run validation and comparison between the existing scenario file and a newly selected CSV file."""
    if file_type not in FILE_MAP or file_type == "price_overrides":
        raise HTTPException(status_code=400, detail=f"Invalid CSV file type: {file_type}")

    scen_dir = SCENARIOS_DIR / name
    if not scen_dir.exists():
        raise HTTPException(status_code=404, detail=f"Scenario '{name}' does not exist.")

    contents = await file.read()
    string_data = contents.decode("utf-8")
    io_data = StringIO(string_data)

    try:
        # 1. Parse and validate the new file
        df_new = None
        if file_type == "volume_data":
            df_new = load_planning_volume_data(io_data)
        elif file_type == "base_prices":
            df_new = load_pricing_data(io_data)
        elif file_type == "base_costs":
            df_new = load_cost_data(io_data)
        elif file_type == "base_var_costs":
            df_new = load_variable_cost_data(io_data)
        elif file_type == "base_dist_costs":
            df_new = load_distribution_cost_data(io_data)
        elif file_type == "fx_rates":
            df_new = load_fx_data(io_data)
        elif file_type == "plant_currency":
            df_new = load_plant_currency_data(io_data)

        # 2. Get keys and value columns for comparison
        configs_dict = {
            "volume_data": (["Material ID", "Plant", "Sold to ID", "Ship to ID", "Date"], ["Volume"]),
            "base_prices": (["Sold to ID", "Ship to ID", "Material ID"], ["Price"]),
            "base_costs": (["Plant", "Material ID", "Date"], ["Cost"]),
            "base_var_costs": (["Material ID"], ["Variable Cost"]),
            "base_dist_costs": (["Ship to ID"], ["Distribution Cost"]),
            "fx_rates": (["Period", "Currency"], ["Rate"]),
            "plant_currency": (["Plant"], ["Currency"]),
        }

        keys, val_cols = configs_dict[file_type]

        # 3. Load the existing file
        scen_data = load_scenario_data(name)
        df_base = scen_data.get(file_type)

        # 4. Generate diff logs
        diff_logs = diff_dataframe(df_base, df_new, keys, val_cols)

        # If base is None, it means it's a completely new file upload
        is_new = df_base is None

        # 5. For volume_data: compute per-material volume delta (MT)
        volume_delta_by_material = None
        if file_type == "volume_data" and df_new is not None:
            # Sum volumes by Material ID in new file
            new_by_mat = (
                df_new.groupby("Material ID")["Volume"]
                .sum()
                .reset_index()
                .rename(columns={"Volume": "new_volume"})
            )
            if df_base is not None:
                base_by_mat = (
                    df_base.groupby("Material ID")["Volume"]
                    .sum()
                    .reset_index()
                    .rename(columns={"Volume": "base_volume"})
                )
                merged = pd.merge(base_by_mat, new_by_mat, on="Material ID", how="outer")
                merged["base_volume"] = merged["base_volume"].fillna(Decimal("0.000"))
                merged["new_volume"] = merged["new_volume"].fillna(Decimal("0.000"))
                merged["delta"] = merged["new_volume"] - merged["base_volume"]
            else:
                merged = new_by_mat.copy()
                merged["base_volume"] = Decimal("0.000")
                merged["delta"] = merged["new_volume"]

            # Convert Decimal to float for JSON serialisation
            volume_delta_by_material = [
                {
                    "material_id": row["Material ID"],
                    "base_volume": float(row["base_volume"]),
                    "new_volume": float(row["new_volume"]),
                    "delta": float(row["delta"]),
                }
                for _, row in merged.iterrows()
                if float(row["delta"]) != 0.0  # only include changed materials
            ]
            # Sort by absolute delta descending
            volume_delta_by_material.sort(key=lambda x: abs(x["delta"]), reverse=True)

        # 6. For base_prices: compute per-material average price delta
        price_delta_by_material = None
        if file_type == "base_prices" and df_new is not None:
            new_avg = (
                df_new.groupby("Material ID")["Price"]
                .apply(lambda s: s.sum() / Decimal(len(s)) if len(s) > 0 else Decimal("0.000"))
                .reset_index()
                .rename(columns={"Price": "new_avg_price"})
            )
            if df_base is not None:
                base_avg = (
                    df_base.groupby("Material ID")["Price"]
                    .apply(lambda s: s.sum() / Decimal(len(s)) if len(s) > 0 else Decimal("0.000"))
                    .reset_index()
                    .rename(columns={"Price": "base_avg_price"})
                )
                merged_p = pd.merge(base_avg, new_avg, on="Material ID", how="outer")
                merged_p["base_avg_price"] = merged_p["base_avg_price"].fillna(Decimal("0.000"))
                merged_p["new_avg_price"] = merged_p["new_avg_price"].fillna(Decimal("0.000"))
                merged_p["delta"] = merged_p["new_avg_price"] - merged_p["base_avg_price"]
            else:
                merged_p = new_avg.copy()
                merged_p["base_avg_price"] = Decimal("0.000")
                merged_p["delta"] = merged_p["new_avg_price"]

            price_delta_by_material = [
                {
                    "material_id": row["Material ID"],
                    "base_avg_price": float(row["base_avg_price"]),
                    "new_avg_price": float(row["new_avg_price"]),
                    "delta": float(row["delta"]),
                    "pct_change": (
                        round(float(row["delta"]) / float(row["base_avg_price"]) * 100, 2)
                        if float(row["base_avg_price"]) != 0.0 else None
                    ),
                }
                for _, row in merged_p.iterrows()
                if float(row["delta"]) != 0.0
            ]
            # Sort by absolute delta descending
            price_delta_by_material.sort(key=lambda x: abs(x["delta"]), reverse=True)

        # 7. For base_var_costs: compute per-material variable cost delta
        var_cost_delta_by_material = None
        if file_type == "base_var_costs" and df_new is not None:
            new_v = (
                df_new.groupby("Material ID")["Variable Cost"]
                .apply(lambda s: s.sum() / Decimal(len(s)) if len(s) > 0 else Decimal("0.000"))
                .reset_index()
                .rename(columns={"Variable Cost": "new_var_cost"})
            )
            if df_base is not None:
                base_v = (
                    df_base.groupby("Material ID")["Variable Cost"]
                    .apply(lambda s: s.sum() / Decimal(len(s)) if len(s) > 0 else Decimal("0.000"))
                    .reset_index()
                    .rename(columns={"Variable Cost": "base_var_cost"})
                )
                merged_v = pd.merge(base_v, new_v, on="Material ID", how="outer")
                merged_v["base_var_cost"] = merged_v["base_var_cost"].fillna(Decimal("0.000"))
                merged_v["new_var_cost"] = merged_v["new_var_cost"].fillna(Decimal("0.000"))
                merged_v["delta"] = merged_v["new_var_cost"] - merged_v["base_var_cost"]
            else:
                merged_v = new_v.copy()
                merged_v["base_var_cost"] = Decimal("0.000")
                merged_v["delta"] = merged_v["new_var_cost"]

            var_cost_delta_by_material = [
                {
                    "material_id": row["Material ID"],
                    "base_var_cost": float(row["base_var_cost"]),
                    "new_var_cost": float(row["new_var_cost"]),
                    "delta": float(row["delta"]),
                }
                for _, row in merged_v.iterrows()
                if float(row["delta"]) != 0.0
            ]
            var_cost_delta_by_material.sort(key=lambda x: abs(x["delta"]), reverse=True)

        # 8. For fx_rates: compute per-currency average rate delta
        fx_delta_by_currency = None
        if file_type == "fx_rates" and df_new is not None:
            new_f = (
                df_new.groupby("Currency")["Rate"]
                .apply(lambda s: s.sum() / Decimal(len(s)) if len(s) > 0 else Decimal("0.0000"))
                .reset_index()
                .rename(columns={"Rate": "new_rate"})
            )
            if df_base is not None:
                base_f = (
                    df_base.groupby("Currency")["Rate"]
                    .apply(lambda s: s.sum() / Decimal(len(s)) if len(s) > 0 else Decimal("0.0000"))
                    .reset_index()
                    .rename(columns={"Rate": "base_rate"})
                )
                merged_f = pd.merge(base_f, new_f, on="Currency", how="outer")
                merged_f["base_rate"] = merged_f["base_rate"].fillna(Decimal("0.0000"))
                merged_f["new_rate"] = merged_f["new_rate"].fillna(Decimal("0.0000"))
                merged_f["delta"] = merged_f["new_rate"] - merged_f["base_rate"]
            else:
                merged_f = new_f.copy()
                merged_f["base_rate"] = Decimal("0.0000")
                merged_f["delta"] = merged_f["new_rate"]

            fx_delta_by_currency = [
                {
                    "currency": row["Currency"],
                    "base_rate": float(row["base_rate"]),
                    "new_rate": float(row["new_rate"]),
                    "delta": float(row["delta"]),
                }
                for _, row in merged_f.iterrows()
                if float(row["delta"]) != 0.0
            ]
            fx_delta_by_currency.sort(key=lambda x: abs(x["delta"]), reverse=True)

        return {
            "status": "success",
            "is_new": is_new,
            "diff": diff_logs,
            "volume_delta_by_material": volume_delta_by_material,
            "price_delta_by_material": price_delta_by_material,
            "var_cost_delta_by_material": var_cost_delta_by_material,
            "fx_delta_by_currency": fx_delta_by_currency,
        }
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Validation failed: {str(e)}")


@router.post("/scenarios/{name}/save")
def trigger_scenario_save(name: str):
    """Recalculate auto-change logs and save scenario state."""
    try:
        data = load_scenario_data(name)
        save_scenario_data(name, data)
        return {"status": "success", "message": f"Scenario '{name}' saved successfully."}
    except ValueError as e:
        raise HTTPException(status_code=404, detail=str(e))
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@router.post("/scenarios/{name}/calculate", response_model=CalculationResponse)
def calculate_scenario(name: str):
    """Run simulation pipeline and return USD/LC aggregates and calculated preview."""
    try:
        data = load_scenario_data(name)
        
        # Verify all input datasets are loaded
        missing = [k for k, v in data.items() if v is None and k != "price_overrides"]
        if missing:
            raise HTTPException(status_code=400, detail=f"Scenario has missing files: {missing}")

        # Convert price overrides dict list to PriceOverride classes
        price_overrides_list = []
        for item in data["price_overrides"]:
            price_overrides_list.append(PriceOverride(
                material_id=item["Material ID"],
                sold_to_id=item["Sold to ID"],
                ship_to_id=item["Ship to ID"],
                start_month=item["Date"],
                new_price=item["Price"]
            ))

        resolved_prices = resolve_monthly_prices(data["base_prices"], price_overrides_list)
        
        df_calc = run_simulation_pipeline(
            volume_df=data["volume_data"],
            resolved_prices=resolved_prices,
            base_costs=data["base_costs"],
            base_var_costs=data["base_var_costs"],
            base_dist_costs=data["base_dist_costs"],
            plant_currency_mapping_df=data["plant_currency"],
            fx_rates_df=data["fx_rates"],
        )

        metrics = generate_summary_metrics(df_calc)
        preview_records = df_to_records(df_calc.head(100))

        return {
            "metrics": metrics,
            "preview": preview_records
        }
    except HTTPException as e:
        raise e
    except ValueError as e:
        raise HTTPException(status_code=400, detail=f"Calculation error: {str(e)}")
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Internal Server Error: {str(e)}")


@router.get("/scenarios/{name}/export")
def export_scenario_sac(name: str):
    """Export the SAC-ready CSV for the scenario."""
    try:
        data = load_scenario_data(name)
        missing = [k for k, v in data.items() if v is None and k != "price_overrides"]
        if missing:
            raise HTTPException(status_code=400, detail=f"Scenario has missing files: {missing}")

        price_overrides_list = []
        for item in data["price_overrides"]:
            price_overrides_list.append(PriceOverride(
                material_id=item["Material ID"],
                sold_to_id=item["Sold to ID"],
                ship_to_id=item["Ship to ID"],
                start_month=item["Date"],
                new_price=item["Price"]
            ))

        resolved_prices = resolve_monthly_prices(data["base_prices"], price_overrides_list)
        
        df_calc = run_simulation_pipeline(
            volume_df=data["volume_data"],
            resolved_prices=resolved_prices,
            base_costs=data["base_costs"],
            base_var_costs=data["base_var_costs"],
            base_dist_costs=data["base_dist_costs"],
            plant_currency_mapping_df=data["plant_currency"],
            fx_rates_df=data["fx_rates"],
        )

        # Write to memory buffer
        stream = io.StringIO()
        df_calc.to_csv(stream, index=False)
        response = StreamingResponse(
            io.BytesIO(stream.getvalue().encode("utf-8")),
            media_type="text/csv",
        )
        response.headers["Content-Disposition"] = f"attachment; filename=calculated_revenue_and_rm_cost_{name}.csv"
        return response
    except HTTPException as e:
        raise e
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Export failed: {str(e)}")


@router.post("/compare", response_model=ScenarioDifference)
def compare_two_scenarios(payload: CompareRequest):
    """Compare two scenarios side-by-side, returning metrics deltas and VCM variances."""
    scen_a = payload.scenario_a
    scen_b = payload.scenario_b

    try:
        data_a = load_scenario_data(scen_a)
        data_b = load_scenario_data(scen_b)

        missing_a = [k for k, v in data_a.items() if v is None and k != "price_overrides"]
        missing_b = [k for k, v in data_b.items() if v is None and k != "price_overrides"]

        if missing_a:
            raise HTTPException(status_code=400, detail=f"Scenario A is incomplete: {missing_a}")
        if missing_b:
            raise HTTPException(status_code=400, detail=f"Scenario B is incomplete: {missing_b}")

        # Run Scenario A
        over_list_a = [PriceOverride(item["Material ID"], item["Sold to ID"], item["Ship to ID"], item["Date"], item["Price"]) for item in data_a["price_overrides"]]
        prices_res_a = resolve_monthly_prices(data_a["base_prices"], over_list_a)
        df_calc_a = run_simulation_pipeline(
            data_a["volume_data"], prices_res_a, data_a["base_costs"],
            data_a["base_var_costs"], data_a["base_dist_costs"],
            data_a["plant_currency"], data_a["fx_rates"]
        )
        metrics_a = generate_summary_metrics(df_calc_a)

        # Run Scenario B
        over_list_b = [PriceOverride(item["Material ID"], item["Sold to ID"], item["Ship to ID"], item["Date"], item["Price"]) for item in data_b["price_overrides"]]
        prices_res_b = resolve_monthly_prices(data_b["base_prices"], over_list_b)
        df_calc_b = run_simulation_pipeline(
            data_b["volume_data"], prices_res_b, data_b["base_costs"],
            data_b["base_var_costs"], data_b["base_dist_costs"],
            data_b["plant_currency"], data_b["fx_rates"]
        )
        metrics_b = generate_summary_metrics(df_calc_b)

        # Absolute Delta
        abs_diff = {}
        compare_keys = [
            "total_volume",
            "total_revenue_usd",
            "total_rm_cost_usd",
            "total_var_cost_usd",
            "total_dist_cost_usd",
            "total_vcm_usd",
            "weighted_avg_price_usd",
            "weighted_avg_vcm_usd",
        ]
        for key in compare_keys:
            abs_diff[key] = metrics_b[key] - metrics_a[key]

        # Percentage Delta
        pct_diff = {}
        for key in compare_keys:
            val_a = metrics_a[key]
            delta = metrics_b[key] - val_a
            pct_diff[key] = f"{(delta / val_a * 100):+.1f}%" if val_a != 0 else "0.0%"

        # Plant VCM Variance
        plant_a = df_calc_a.groupby("Plant")[["Volume", "VCM_USD"]].sum().reset_index()
        plant_b = df_calc_b.groupby("Plant")[["Volume", "VCM_USD"]].sum().reset_index()
        plant_merged = pd.merge(plant_a, plant_b, on="Plant", suffixes=("_A", "_B"), how="outer").fillna(0)
        plant_merged["Volume Delta"] = plant_merged["Volume_B"] - plant_merged["Volume_A"]
        plant_merged["VCM Delta"] = plant_merged["VCM_USD_B"] - plant_merged["VCM_USD_A"]
        plant_variance = df_to_records(plant_merged)

        # Material VCM Variance
        mat_a = df_calc_a.groupby(["Material", "Material ID"])[["Volume", "VCM_USD"]].sum().reset_index()
        mat_b = df_calc_b.groupby(["Material", "Material ID"])[["Volume", "VCM_USD"]].sum().reset_index()
        mat_merged = pd.merge(mat_a, mat_b, on=["Material", "Material ID"], suffixes=("_A", "_B"), how="outer").fillna(0)
        mat_merged["Volume Delta"] = mat_merged["Volume_B"] - mat_merged["Volume_A"]
        mat_merged["VCM Delta"] = mat_merged["VCM_USD_B"] - mat_merged["VCM_USD_A"]
        mat_variance = df_to_records(mat_merged)

        # Diff logs
        diff_logs = diff_scenarios_inputs(scen_a, scen_b)

        return {
            "absolute_diff": abs_diff,
            "percentage_diff": pct_diff,
            "change_log": diff_logs,
            "vcm_variance_by_plant": plant_variance,
            "vcm_variance_by_material": mat_variance,
        }
    except HTTPException as e:
        raise e
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Comparison failed: {str(e)}")


@router.post("/compare/bridge", response_model=BridgeResponse)
def compare_bridge(payload: CompareRequest):
    """Calculate PVM + FX margin bridge between two scenarios."""
    scen_a = payload.scenario_a
    scen_b = payload.scenario_b

    try:
        data_a = load_scenario_data(scen_a)
        data_b = load_scenario_data(scen_b)

        missing_a = [k for k, v in data_a.items() if v is None and k != "price_overrides"]
        missing_b = [k for k, v in data_b.items() if v is None and k != "price_overrides"]

        if missing_a:
            raise HTTPException(status_code=400, detail=f"Scenario A is incomplete: {missing_a}")
        if missing_b:
            raise HTTPException(status_code=400, detail=f"Scenario B is incomplete: {missing_b}")

        # Run Scenario A
        over_list_a = [PriceOverride(item["Material ID"], item["Sold to ID"], item["Ship to ID"], item["Date"], item["Price"]) for item in data_a["price_overrides"]]
        prices_res_a = resolve_monthly_prices(data_a["base_prices"], over_list_a)
        df_calc_a = run_simulation_pipeline(
            data_a["volume_data"], prices_res_a, data_a["base_costs"],
            data_a["base_var_costs"], data_a["base_dist_costs"],
            data_a["plant_currency"], data_a["fx_rates"]
        )

        # Run Scenario B
        over_list_b = [PriceOverride(item["Material ID"], item["Sold to ID"], item["Ship to ID"], item["Date"], item["Price"]) for item in data_b["price_overrides"]]
        prices_res_b = resolve_monthly_prices(data_b["base_prices"], over_list_b)
        df_calc_b = run_simulation_pipeline(
            data_b["volume_data"], prices_res_b, data_b["base_costs"],
            data_b["base_var_costs"], data_b["base_dist_costs"],
            data_b["plant_currency"], data_b["fx_rates"]
        )

        # Run row-by-row bridge calculations
        df_bridge = calculate_margin_bridge(
            df_calc_a,
            df_calc_b,
            data_a["fx_rates"],
            data_b["fx_rates"]
        )

        # Aggregate Summary
        summary = summarize_margin_bridge(df_bridge)

        # Breakdown by Material
        df_mat = (
            df_bridge.groupby(["Material", "Material ID"])[
                ["Vol_A", "Vol_B", "VCM_USD_A", "VCM_USD_B", "Volume_Effect", "Price_Effect", "Cost_Effect", "FX_Effect"]
            ]
            .sum()
            .reset_index()
        )
        # Sort by VCM Delta absolute value descending
        df_mat["delta"] = df_mat["VCM_USD_B"] - df_mat["VCM_USD_A"]
        df_mat["abs_delta"] = df_mat["delta"].abs()
        df_mat.sort_values(by="abs_delta", ascending=False, inplace=True)
        df_mat.drop(columns=["abs_delta", "delta"], inplace=True)

        # Breakdown by Month
        df_month = (
            df_bridge.groupby(["Date"])[
                ["Vol_A", "Vol_B", "VCM_USD_A", "VCM_USD_B", "Volume_Effect", "Price_Effect", "Cost_Effect", "FX_Effect"]
            ]
            .sum()
            .reset_index()
        )
        df_month["sort_key"] = df_month["Date"].apply(lambda x: str(x))
        df_month.sort_values(by="sort_key", inplace=True)
        df_month.drop(columns=["sort_key"], inplace=True)

        by_mat_records = df_to_records(df_mat)
        by_month_records = df_to_records(df_month)
        commentary_list = generate_bridge_commentary(summary, by_mat_records, by_month_records)

        return {
            "summary": summary,
            "by_material": by_mat_records,
            "by_month": by_month_records,
            "raw_preview": df_to_records(df_bridge),
            "commentary": commentary_list,
        }
    except HTTPException as e:
        raise e
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Bridge calculation failed: {str(e)}")


@router.post("/login-log")
def log_user_login(payload: LoginLogInput):
    """Log user access to the backend system for audit purposes."""
    try:
        import datetime
        from pathlib import Path
        
        log_file = LOGIN_AUDIT_LOG
        log_file.parent.mkdir(parents=True, exist_ok=True)
        
        timestamp = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        log_line = f"[{timestamp}] User: {payload.name} ({payload.email}) signed in via {payload.provider}\n"
        
        with open(log_file, "a", encoding="utf-8") as f:
            f.write(log_line)
            
        return {"status": "success", "message": "Login logged successfully."}
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


TEMPLATE_DATA = {
    "volume_data": {
        "columns": ["Material", "Material ID", "Date", "Sold to ID", "Sold to", "Ship to ID", "Ship to", "Plant", "Volume"],
        "rows": [
            ["Material A", "MAT-A", "2026-01", "CUST-01", "Customer 01", "SHIP-01", "Ship to 01", "PLANT-01", 10.000],
            ["Material A", "MAT-A", "2026-02", "CUST-01", "Customer 01", "SHIP-01", "Ship to 01", "PLANT-01", 15.000],
            ["Material B", "MAT-B", "2026-01", "CUST-02", "Customer 02", "SHIP-02", "Ship to 02", "PLANT-01", 5.500],
        ]
    },
    "base_prices": {
        "columns": ["Sold to ID", "Ship to ID", "Material ID", "Price"],
        "rows": [
            ["CUST-01", "SHIP-01", "MAT-A", 150.00],
            ["CUST-02", "SHIP-02", "MAT-B", 230.50],
            ["CUST-01", "SHIP-01", "MAT-B", 180.00],
        ]
    },
    "base_costs": {
        "columns": ["Plant", "Material ID", "Period", "Cost"],
        "rows": [
            ["PLANT-01", "MAT-A", "2026-01", 50.00],
            ["PLANT-01", "MAT-A", "2026-02", 52.50],
            ["PLANT-01", "MAT-B", "2026-01", 65.00],
        ]
    },
    "base_var_costs": {
        "columns": ["Material", "Material ID", "Variable Cost"],
        "rows": [
            ["Material A", "MAT-A", 12.50],
            ["Material B", "MAT-B", 18.00],
            ["Material C", "MAT-C", 14.20],
        ]
    },
    "base_dist_costs": {
        "columns": ["Ship to", "Ship to ID", "Distribution Cost"],
        "rows": [
            ["Ship to 01", "SHIP-01", 5.00],
            ["Ship to 02", "SHIP-02", 7.20],
            ["Ship to 03", "SHIP-03", 4.50],
        ]
    },
    "fx_rates": {
        "columns": ["Period", "Currency", "Rate"],
        "rows": [
            ["2026-01", "EUR", 0.9200],
            ["2026-02", "EUR", 0.9150],
            ["2026-01", "CAD", 1.3700],
        ]
    },
    "plant_currency": {
        "columns": ["Plant", "Currency"],
        "rows": [
            ["PLANT-01", "USD"],
            ["PLANT-02", "EUR"],
            ["PLANT-03", "CAD"],
        ]
    }
}


@router.get("/templates/{file_type}")
def get_file_template(file_type: str):
    """Generate and return an Excel template populated with headers and example rows."""
    if file_type not in TEMPLATE_DATA:
        raise HTTPException(status_code=400, detail="Invalid template file type requested.")
    
    data = TEMPLATE_DATA[file_type]
    df = pd.DataFrame(data["rows"], columns=data["columns"])
    
    # Write to Excel in memory using openpyxl
    out = io.BytesIO()
    with pd.ExcelWriter(out, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Template")
    
    out.seek(0)
    
    headers = {
        "Content-Disposition": f'attachment; filename="{file_type}_template.xlsx"'
    }
    
    return StreamingResponse(
        out,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers=headers
    )


@router.get("/admin/login-logs")
def get_login_logs(email: str):
    """Retrieve user login audit logs. Restricted to administrator."""
    admin_emails_env = os.getenv("ADMIN_EMAILS", "")
    if admin_emails_env:
        authorized_emails = [e.strip() for e in admin_emails_env.split(",") if e.strip()]
    else:
        authorized_emails = ["ney.morais@gmail.com", "ney.morais@outlook.com"]

    if email not in authorized_emails:
        raise HTTPException(
            status_code=403,
            detail="Forbidden: Unauthorized user email."
        )
    
    log_file = LOGIN_AUDIT_LOG
    if not log_file.exists():
        return []
        
    import re
    log_pattern = re.compile(r"^\[(.*?)\] User: (.*?) \((.*?)\) signed in via (.*?)$")
    
    logs = []
    try:
        with open(log_file, "r", encoding="utf-8") as f:
            for line in f:
                line = line.strip()
                if not line:
                    continue
                match = log_pattern.match(line)
                if match:
                    timestamp, name, user_email, provider = match.groups()
                    logs.append({
                        "timestamp": timestamp,
                        "name": name,
                        "email": user_email,
                        "provider": provider
                    })
        # Return newest logs first
        logs.reverse()
        return logs
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@router.post("/export/bridge-pptx")
def export_bridge_pptx(payload: PPTXExportInput):
    """Generate a high-fidelity PowerPoint slide with a matplotlib-rendered waterfall chart."""
    try:
        import matplotlib
        matplotlib.use("Agg")  # headless backend
        import matplotlib.pyplot as plt
        import matplotlib.patches as mpatches
        from matplotlib.patches import FancyBboxPatch
        import numpy as np

        from pptx import Presentation
        from pptx.util import Inches, Pt, Emu
        from pptx.dml.color import RGBColor
        from pptx.enum.text import PP_ALIGN

        # ── Theme colors (matching the app's CSS) ──
        is_dark = payload.theme != "light"

        if is_dark:
            bg_color = "#0f1225"
            primary_color = "#22b8cf"
            positive_color = "#10b981"
            negative_color = "#ef4444"
            muted_color = "#8b95a8"
            foreground_color = "#eef0f4"
            border_color = "#2a2f4a"
            connector_color = "#4a5068"
            slide_bg_rgb = RGBColor(15, 18, 37)
            title_rgb = RGBColor(238, 240, 244)
            subtitle_rgb = RGBColor(139, 149, 168)
            footer_rgb = RGBColor(100, 108, 130)
        else:
            bg_color = "#f8f9fc"
            primary_color = "#2563b0"
            positive_color = "#16a34a"
            negative_color = "#dc2626"
            muted_color = "#64748b"
            foreground_color = "#1e293b"
            border_color = "#e2e5ea"
            connector_color = "#94a3b8"
            slide_bg_rgb = RGBColor(248, 249, 252)
            title_rgb = RGBColor(30, 41, 59)
            subtitle_rgb = RGBColor(100, 116, 139)
            footer_rgb = RGBColor(148, 163, 184)

        # ── Data preparation ──
        base = float(payload.vcm_usd_a)
        vol = float(payload.volume_effect)
        price = float(payload.price_effect)
        cost = float(payload.cost_effect)
        fx = float(payload.fx_effect)
        target = float(payload.vcm_usd_b)

        labels = [
            payload.scenario_a,
            "Volume\nEffect",
            "Price\nEffect",
            "Cost\nEffect",
            "FX\nEffect",
            payload.scenario_b,
        ]

        # Running cumulative positions
        step1 = base
        step2 = step1 + vol
        step3 = step2 + price
        step4 = step3 + cost
        step5 = step4 + fx

        steps = [
            {"bottom": 0, "height": base, "change": base, "type": "base"},
            {"bottom": min(step1, step2), "height": abs(vol), "change": vol, "type": "var"},
            {"bottom": min(step2, step3), "height": abs(price), "change": price, "type": "var"},
            {"bottom": min(step3, step4), "height": abs(cost), "change": cost, "type": "var"},
            {"bottom": min(step4, step5), "height": abs(fx), "change": fx, "type": "var"},
            {"bottom": 0, "height": target, "change": target, "type": "target"},
        ]

        # Connection line y-values (the "end" of each bar, before the next)
        end_values = [base, step2, step3, step4, step5]

        # ── Currency formatter ──
        def fmt_curr(val: float) -> str:
            sign = "+" if val > 0 else ""
            av = abs(val)
            if av >= 1_000_000:
                return f"{sign}${val / 1_000_000:,.1f}M"
            elif av >= 1_000:
                return f"{sign}${val / 1_000:,.1f}k"
            else:
                return f"{sign}${val:,.0f}"

        # ── Matplotlib chart rendering ──
        fig, ax = plt.subplots(figsize=(12, 5.5))
        fig.patch.set_facecolor(bg_color)
        ax.set_facecolor(bg_color)

        n = len(steps)
        bar_width = 0.55
        x_positions = np.arange(n)

        for i, step in enumerate(steps):
            # Determine bar color
            if step["type"] in ("base", "target"):
                color = primary_color
            elif step["change"] > 0:
                color = positive_color
            elif step["change"] < 0:
                color = negative_color
            else:
                color = muted_color

            # Draw rounded bar using FancyBboxPatch
            bar_rect = FancyBboxPatch(
                (x_positions[i] - bar_width / 2, step["bottom"]),
                bar_width,
                max(step["height"], abs(step["change"]) * 0.003 or 1),  # min visible height
                boxstyle="round,pad=0,rounding_size=0.06",
                facecolor=color,
                edgecolor="none",
                zorder=3,
                alpha=0.92,
            )
            ax.add_patch(bar_rect)

            # Value label above or below the bar
            val_text = fmt_curr(step["change"])
            if step["type"] in ("base", "target"):
                label_color = foreground_color
            elif step["change"] > 0:
                label_color = positive_color
            elif step["change"] < 0:
                label_color = negative_color
            else:
                label_color = muted_color

            bar_top = step["bottom"] + step["height"]
            if step["change"] >= 0:
                label_y = bar_top
                va = "bottom"
            else:
                label_y = step["bottom"]
                va = "top"

            offset = (ax.get_ylim()[1] - ax.get_ylim()[0]) * 0.02 if ax.get_ylim()[1] != ax.get_ylim()[0] else 1000
            label_y_adj = label_y + offset if va == "bottom" else label_y - offset

            ax.text(
                x_positions[i], label_y_adj, val_text,
                ha="center", va=va,
                fontsize=9.5, fontweight="bold", color=label_color,
                fontfamily="sans-serif",
                zorder=5,
            )

        # ── Dashed connector lines between bars ──
        for i in range(n - 1):
            y_line = end_values[i]
            x_start = x_positions[i] + bar_width / 2
            x_end = x_positions[i + 1] - bar_width / 2
            ax.plot(
                [x_start, x_end], [y_line, y_line],
                color=connector_color,
                linewidth=1.0,
                linestyle=(0, (4, 3)),  # dashed
                zorder=2,
            )

        # ── Axis styling ──
        ax.set_xticks(x_positions)
        ax.set_xticklabels(labels, fontsize=9.5, fontweight="semibold",
                           color=muted_color, fontfamily="sans-serif")
        ax.tick_params(axis="x", length=0, pad=10)

        # Y-axis: minimal, no ticks, no labels
        ax.tick_params(axis="y", left=False, labelleft=False)

        # Remove spines
        for spine in ax.spines.values():
            spine.set_visible(False)

        # Light horizontal zero line
        ax.axhline(y=0, color=border_color, linewidth=0.8, zorder=1)

        # Auto y-limits with padding
        all_vals = [s["bottom"] for s in steps] + [s["bottom"] + s["height"] for s in steps]
        y_min = min(all_vals)
        y_max = max(all_vals)
        y_pad = (y_max - y_min) * 0.18 if y_max != y_min else 1000
        ax.set_ylim(y_min - y_pad, y_max + y_pad)
        ax.set_xlim(-0.6, n - 0.4)

        # ── Legend ──
        legend_elements = [
            mpatches.Patch(facecolor=primary_color, edgecolor="none", label="Scenario Totals"),
            mpatches.Patch(facecolor=positive_color, edgecolor="none", label="Positive Impact (+)"),
            mpatches.Patch(facecolor=negative_color, edgecolor="none", label="Negative Impact (−)"),
        ]
        leg = ax.legend(
            handles=legend_elements,
            loc="lower center",
            bbox_to_anchor=(0.5, -0.18),
            ncol=3,
            frameon=False,
            fontsize=9,
            labelcolor=foreground_color,
            handlelength=1.2,
            handletextpad=0.5,
            columnspacing=2.5,
        )
        for text in leg.get_texts():
            text.set_fontfamily("sans-serif")

        plt.tight_layout(pad=1.0)

        # ── Export chart as high-res PNG to memory ──
        chart_buf = io.BytesIO()
        fig.savefig(chart_buf, format="png", dpi=300, bbox_inches="tight",
                    facecolor=bg_color, edgecolor="none", pad_inches=0.3)
        plt.close(fig)
        chart_buf.seek(0)

        # ── Build PPTX slide ──
        prs = Presentation()
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.5)

        blank_layout = prs.slide_layouts[6]
        slide = prs.slides.add_slide(blank_layout)

        # Slide background
        bg = slide.background
        fill = bg.fill
        fill.solid()
        fill.fore_color.rgb = slide_bg_rgb

        # ── Title text (top-left) ──
        title_box = slide.shapes.add_textbox(Inches(0.6), Inches(0.35), Inches(6.0), Inches(1.0))
        tf = title_box.text_frame
        tf.word_wrap = True

        p_sub = tf.paragraphs[0]
        p_sub.text = "VCM WATERFALL BRIDGE"
        p_sub.font.size = Pt(10)
        p_sub.font.bold = True
        p_sub.font.color.rgb = subtitle_rgb
        p_sub.font.name = "Segoe UI"

        p_title = tf.add_paragraph()
        p_title.text = "Scenario VCM Bridge (USD)"
        p_title.font.size = Pt(22)
        p_title.font.bold = True
        p_title.font.color.rgb = title_rgb
        p_title.font.name = "Segoe UI"
        p_title.space_before = Pt(4)

        # ── Filter info (top-right) ──
        filter_box = slide.shapes.add_textbox(Inches(8.5), Inches(0.35), Inches(4.3), Inches(1.0))
        tf_f = filter_box.text_frame
        tf_f.word_wrap = True

        p_mat = tf_f.paragraphs[0]
        p_mat.alignment = PP_ALIGN.RIGHT
        p_mat.text = f"MATERIAL: {payload.material_filter}"
        p_mat.font.size = Pt(9)
        p_mat.font.bold = True
        p_mat.font.color.rgb = subtitle_rgb
        p_mat.font.name = "Segoe UI"

        p_reg = tf_f.add_paragraph()
        p_reg.alignment = PP_ALIGN.RIGHT
        p_reg.text = f"REGION: {payload.region_filter}"
        p_reg.font.size = Pt(9)
        p_reg.font.bold = True
        p_reg.font.color.rgb = subtitle_rgb
        p_reg.font.name = "Segoe UI"
        p_reg.space_before = Pt(3)

        # ── Embed the chart image ──
        if payload.commentary:
            chart_width = Inches(8.5)
            slide.shapes.add_picture(chart_buf, Inches(0.3), Inches(1.5), width=chart_width)

            # Add commentary box on the right
            comm_box = slide.shapes.add_textbox(Inches(9.0), Inches(1.5), Inches(3.8), Inches(5.0))
            tf_c = comm_box.text_frame
            tf_c.word_wrap = True

            # Title
            p_c_title = tf_c.paragraphs[0]
            p_c_title.text = "BRIDGE DRIVER ANALYSIS"
            p_c_title.font.size = Pt(11)
            p_c_title.font.bold = True
            p_c_title.font.color.rgb = title_rgb
            p_c_title.font.name = "Segoe UI"
            p_c_title.space_after = Pt(10)

            # Bullets (excluding source markers and capped at 5)
            comments_to_render = list(payload.commentary)
            if comments_to_render and (comments_to_render[0].startswith("[AI") or comments_to_render[0].startswith("[Det")):
                comments_to_render = comments_to_render[1:]

            comments_to_render = comments_to_render[:5]

            for comment in comments_to_render:
                clean_comment = comment.replace("**", "")
                p_bullet = tf_c.add_paragraph()
                p_bullet.text = f"• {clean_comment}"
                p_bullet.font.size = Pt(9.5)
                p_bullet.font.color.rgb = subtitle_rgb
                p_bullet.font.name = "Segoe UI"
                p_bullet.space_after = Pt(8)
                p_bullet.line_spacing = 1.15
        else:
            chart_left = Inches(0.3)
            chart_top = Inches(1.5)
            chart_width = Inches(12.7)
            slide.shapes.add_picture(chart_buf, chart_left, chart_top, width=chart_width)

        # ── Footer branding ──
        footer_box = slide.shapes.add_textbox(Inches(0.6), Inches(7.0), Inches(5.0), Inches(0.35))
        f_tf = footer_box.text_frame
        f_p = f_tf.paragraphs[0]
        f_p.text = "Financial Planner — FP&A 2026"
        f_p.font.size = Pt(8)
        f_p.font.italic = True
        f_p.font.color.rgb = footer_rgb
        f_p.font.name = "Segoe UI"

        # ── Stream result ──
        out = io.BytesIO()
        prs.save(out)
        out.seek(0)

        headers = {
            "Content-Disposition": 'attachment; filename="vcm_waterfall_bridge.pptx"'
        }

        return StreamingResponse(
            out,
            media_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
            headers=headers
        )
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Failed to generate PPTX: {str(e)}")




